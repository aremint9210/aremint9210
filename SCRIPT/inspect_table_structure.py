from docx import Document
import sys

def inspect_table(file_path):
    doc = Document(file_path)
    if not doc.tables:
        print("No tables found.")
        return
    table = doc.tables[0]
    print(f"Table has {len(table.rows)} rows and {len(table.columns)} columns.")
    for i, cell in enumerate(table.rows[0].cells):
        print(f"Header {i}: {cell.text}")

if __name__ == "__main__":
    inspect_table(sys.argv[1])
