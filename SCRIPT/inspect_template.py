from docx import Document
import sys

try:
    path = r'C:\Users\aremi\Desktop\REPORT CBM\PO 42289160 - 33kV THUNDER CYCLE2 2026\GENERATE REPORT\TEMPLATE\00. EXECUTIVE SUMMARY - Copy.docx'
    doc = Document(path)
    if doc.tables:
        header = [cell.text.strip() for cell in doc.tables[0].rows[0].cells]
        print(f"HEADER_START|{','.join(header)}|HEADER_END")
    else:
        print("NO_TABLES")
except Exception as e:
    print(f"ERROR: {e}")
