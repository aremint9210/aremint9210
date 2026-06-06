import pandas as pd
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

# Configuration
SCRIPT_PATH = Path(__file__).resolve()
BASE_DIR = SCRIPT_PATH.parents[2]
MAPPING_DIR = BASE_DIR / "99. Generate Report" / "EXCELL MAPPING"
DEFAULT_CSV = MAPPING_DIR / "Master List.csv"
TEMPLATE_DIR = BASE_DIR / "99. Generate Report" / "TEMPLATE"
TEMPLATE_DOC = TEMPLATE_DIR / "00. EXECUTIVE SUMMARY - RMU.docx"
OUTPUT_DIR = BASE_DIR / "99. Generate Report" / "OUTPUT FILE" / "GROUP"
OUTPUT_FILE = OUTPUT_DIR / "Executive_Summary.docx"

def set_cell_shading(cell, fill_color):
    """Sets the background color of a table cell."""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:val'), 'clear')
    shading_elm.set(qn('w:color'), 'auto')
    shading_elm.set(qn('w:fill'), fill_color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_summary(csv_path=None):
    print("Starting generate_summary...")
    if csv_path:
        csv_file = Path(csv_path)
    else:
        csv_file = DEFAULT_CSV

    if not csv_file.exists():
        print(f"Error: CSV file not found at {csv_file}")
        return

    # Determine Output Directory (Consolidated under INDIVIDUAL/Substation)
    if csv_path:
        csv_prefix = csv_file.stem
        output_dir = BASE_DIR / "99. Generate Report" / "OUTPUT FILE" / "INDIVIDUAL" / csv_prefix
        output_file = output_dir / "Group 0.docx"
    else:
        output_dir = OUTPUT_DIR # Fallback to GROUP
        output_file = output_dir / "Group 0.docx"

    print(f"Reading data from: {csv_file.name}")
    try:
        df = pd.read_csv(csv_file)
        print(f"Successfully read {len(df)} rows.")
    except Exception as e:
        print(f"Error reading CSV: {e}")
        return

    # ... (rest of the template detection logic)

    # Determine Template based on Switchgear Type
    swg_type = str(df.iloc[0].get('{{ swg.type }}', '')).strip()
    swg_manuf = str(df.iloc[0].get('{{ swg.manufacturer }}', '')).strip()
    print(f"Detected Switchgear Type: {swg_type}")
    
    if swg_type == "VCB":
        template_doc = TEMPLATE_DIR / "00. EXECUTIVE SUMMARY - VCB.docx"
    else:
        template_doc = TEMPLATE_DIR / "00. EXECUTIVE SUMMARY - RMU.docx"

    # Open existing template
    if not template_doc.exists():
        print(f"Error: Template not found at {template_doc}")
        return
    
    print(f"Using template: {template_doc}")
    try:
        doc = Document(template_doc)
        print("Template loaded successfully.")
    except Exception as e:
        print(f"Error loading template: {e}")
        return
    
    # Assuming the first table is the one to populate
    if len(doc.tables) > 0:
        table = doc.tables[0]
        # Re-create the table if it doesn't have 5 columns or just clear it
        # Safest is to delete and add new to ensure 5 columns
        parent = table._element.getparent()
        parent.remove(table._element)
    
    table = doc.add_table(rows=1, cols=5)
    try:
        table.style = 'Table Grid'
    except KeyError:
        print("Warning: 'Table Grid' style not found in template. Using default table style.")
    
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Set Headers
    headers = ["NO.", "EQUIPMENT", "DEFECT AREA", "IR (Abs T/ΔT)", "DEFECT"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        # Bold and shading for headers
        p = table.rows[0].cells[i].paragraphs[0]
        run = p.runs[0] if p.runs else p.add_run()
        run.font.bold = True
        set_cell_shading(table.rows[0].cells[i], "D9D9D9") # Light Gray

    # Category Mapping Helper
    def get_category(type_v):
        tv = str(type_v).upper()
        if "OVERVIEW" in tv or "PANEL" in tv:
            return "SWG"
        if "TX1" in tv or "POWER TX1" in tv:
            return "TX1"
        if "TX2" in tv or "POWER TX2" in tv:
            return "TX2"
        if "FP" in tv or "FEEDER PILLAR" in tv:
            return "FP"
        if "BATTERY" in tv:
            return "BATTERY"
        if "BLACK BOX" in tv:
            return "BLACKBOX"
        return tv

    # Logic to group equipment
    processed_combinations = set()
    item_sequence = 0
    last_category = None
    
    # For merging
    last_equip_desc = None
    last_equip_cell = None
    last_no_cell = None

    for _, row in df.iterrows():
        # Clean up identifiers
        raw_no = str(row.get('No', '')).strip()
        if raw_no == 'nan': raw_no = ''
        
        type_val = str(row.get('Type', '')).strip()
        if not type_val or type_val == 'nan':
            continue

        # Determine Category for sequential numbering
        current_cat = get_category(type_val)
        
        # Determine Area
        area_main = str(row.get('{{ swg.area }}', '')).strip()
        if not area_main or area_main == 'nan':
            area_main = str(row.get('{{area}}', '')).strip()
        if not area_main or area_main == 'nan':
            area_main = str(row.get('Attribute', '')).strip()
            
        area = area_main
        if type_val.startswith('TX') or 'POWER TX' in type_val or 'LOCALT TX' in type_val:
            side = str(row.get('{{ panel.name }}', '')).strip()
            if side in ['HV SIDE', 'LV SIDE']:
                prefix = 'HV ' if side == 'HV SIDE' else 'LV '
                if area_main in ['BUSHING', 'CABLE', 'CABLE SPLIT']:
                    area = f"{prefix}{area_main}"

        # Check if we already processed this combo
        combo = (type_val, area)
        if combo in processed_combinations:
            continue

        # Extract Equipment Description
        equip_desc = ""
        
        if "Overview" in type_val:
            if swg_type == "RMU SF6":
                # Match reference PDF naming convention if possible
                equip_desc = "MRMU, INDKOM"
            else:
                equip_desc = f"{swg_type} - {swg_manuf}"
        elif type_val.startswith('TX') and type_val != 'TX' or 'POWER TX' in type_val or 'LOCALT TX' in type_val:
            tx_num = "".join(filter(str.isdigit, type_val))
            if not tx_num: tx_num = "1"
            
            manuf = str(row.get(f'{{{{ tx{tx_num}.manufacturer }}}}', '')).strip()
            r = str(row.get(f'{{{{ tx{tx_num}.rating }}}}', '')).strip()
            if not manuf or manuf == 'nan': manuf = str(row.get('{{ptx1_manufacturer}}', '')).strip()
            if not r or r == 'nan': r = str(row.get('{{ptx1_rating}}', '')).strip()
            
            if r: r = r.lower().replace('kva', 'kVA').replace('  ', ' ')
            equip_desc = f"TX{tx_num} - {manuf} {r}"
        elif 'panel' in type_val.lower():
            p_no = str(row.get('{{ panel.linknumber }}', '')).strip()
            if not p_no or p_no == 'nan': p_no = str(row.get('{{panel_no}}', '')).strip()
            if p_no.endswith('.0'): p_no = p_no.replace('.0', '')
            
            p_name = str(row.get('{{ panel.name }}', '')).strip()
            if not p_name or p_name == 'nan': p_name = str(row.get('{{panel_name}}', '')).strip()
            
            if p_no and p_name and p_no != '-' and p_name != '-':
                equip_desc = f"PANEL {p_no}\n{p_name}"
            elif p_no and p_no != '-':
                equip_desc = f"PANEL {p_no}"
            else:
                equip_desc = p_name if p_name and p_name != '-' else type_val.upper()
        elif 'BATTERY' in type_val.upper():
            b_num = "".join(filter(str.isdigit, type_val))
            if not b_num: b_num = "1"
            equip_desc = f"BATTERY {b_num}"
        else:
            equip_desc = type_val.upper()

        equip_desc = equip_desc.replace('nan', '').replace(' - ', ' ').strip()
        
        # Image exists check
        folder_path = row.get('Folder')
        file_path = row.get('Image File Name')
        image_exists = False
        if pd.notna(folder_path) and pd.notna(file_path):
            img_path = Path(str(folder_path)) / str(file_path)
            if img_path.exists(): image_exists = True
        
        if not image_exists: continue
            
        # Update sequence counter based on Category
        if current_cat != last_category:
            item_sequence += 1
            last_category = current_cat

        # ADD ROW TO TABLE
        row_obj = table.add_row()
        cells = row_obj.cells
        
        cells[2].text = area
        cells[3].text = '-' # IR placeholder
        
        # DEFECT column (index 4)
        if "OVERVIEW" in area.upper():
            cells[4].text = '-'
        else:
            cells[4].text = ''  # Green shading for OK
            set_cell_shading(cells[4], "00B050")
        
        # Equipment and NO. Column merging logic
        if current_cat == "SWG" and item_sequence == 1:
            # Special merging for Switchgear NO column
            if last_no_cell is not None and last_category == "SWG":
                last_no_cell.merge(cells[0])
            else:
                cells[0].text = str(item_sequence)
                last_no_cell = cells[0]
        else:
            # Standard sequential numbering for other categories
            if current_cat == last_category and last_no_cell is not None:
                last_no_cell.merge(cells[0])
            else:
                cells[0].text = str(item_sequence)
                last_no_cell = cells[0]

        # Equipment merging logic
        if equip_desc == last_equip_desc and last_equip_cell is not None:
            last_equip_cell.merge(cells[1])
        else:
            cells[1].text = equip_desc
            last_equip_cell = cells[1]
            last_equip_desc = equip_desc
        
        processed_combinations.add(combo)

    # FINAL PASS: Formatting
    # Total width approx 18cm. 
    # NO: 1cm, EQUIP: 5cm, AREA: 5cm, IR: 3.5cm, DEFECT: 3.5cm
    col_widths = [Cm(1.0), Cm(5.0), Cm(5.0), Cm(3.5), Cm(3.5)]
    
    for row in table.rows:
        row.height = Cm(0.7)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, cell in enumerate(row.cells):
            if i < len(col_widths):
                cell.width = col_widths[i]
                
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                
                for run in paragraph.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10)
                if not paragraph.runs:
                    run = paragraph.add_run()
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10)

    if not output_dir.exists():
        output_dir.mkdir(parents=True)
        
    doc.save(output_file)
    print(f"Summary generated: {output_file}")





if __name__ == "__main__":
    csv_arg = sys.argv[1] if len(sys.argv) > 1 else None
    generate_summary(csv_arg)
